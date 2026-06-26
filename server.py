# -*- coding: utf-8 -*-
"""中医药科学大数据管理平台 - 纯标准库 HTTP 服务器 (零外部依赖)"""
import json
import os
import sys
import csv
import uuid
import urllib.parse
from http.server import HTTPServer, BaseHTTPRequestHandler
from database import init_db, get_db
from services.knowledge_service import KnowledgeService, GraphService, QAService

# 尝试导入 Neo4j 服务函数
NEO4J_AVAILABLE = False
check_health = sync_all_to_neo4j = query_entity_graph = get_statistics = None
try:
    from services.neo4j_service import (
        check_health,
        sync_all_to_neo4j,
        query_entity_graph,
        get_statistics,
    )
    NEO4J_AVAILABLE = True
except ImportError:
    pass

UPLOAD_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "uploads")
TEMPLATE_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "templates")
STATIC_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "static")

MIME_TYPES = {
    ".html": "text/html; charset=utf-8",
    ".css": "text/css; charset=utf-8",
    ".js": "application/javascript; charset=utf-8",
    ".json": "application/json",
    ".png": "image/png",
    ".jpg": "image/jpeg",
    ".svg": "image/svg+xml",
    ".ico": "image/x-icon",
}


class TCMHandler(BaseHTTPRequestHandler):
    """HTTP 请求处理器"""

    def log_message(self, format, *args):
        """简洁日志"""
        sys.stdout.write("[%s] %s\n" % (self.log_date_time_string(), format % args))

    def _send_json(self, data, status=200):
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Access-Control-Allow-Origin", "*")
        self.send_header("Access-Control-Allow-Methods", "GET, POST, PUT, DELETE, OPTIONS")
        self.send_header("Access-Control-Allow-Headers", "Content-Type")
        self.end_headers()
        self.wfile.write(json.dumps(data, ensure_ascii=False).encode("utf-8"))

    def do_OPTIONS(self):
        """处理 CORS 预检请求"""
        self.send_response(200)
        self.send_header("Access-Control-Allow-Origin", "*")
        self.send_header("Access-Control-Allow-Methods", "GET, POST, PUT, DELETE, OPTIONS")
        self.send_header("Access-Control-Allow-Headers", "Content-Type")
        self.end_headers()

    def _send_error_json(self, message, status=400):
        self._send_json({"success": False, "error": message}, status)

    def _read_body(self):
        length = int(self.headers.get("Content-Length", 0))
        return self.rfile.read(length) if length else b""

    def _parse_query(self):
        parsed = urllib.parse.urlparse(self.path)
        return urllib.parse.parse_qs(parsed.query)

    def _get_query_param(self, params, key, default=""):
        vals = params.get(key, [default])
        return vals[0] if vals else default

    # ===== 静态文件 =====
    def _serve_static(self):
        if self.command != "GET":
            return False
        path = self.path
        if path == "/" or path == "":
            return self._serve_index()
        if path.startswith("/static/"):
            return self._serve_file(path)
        if path.startswith("/templates/"):
            return self._serve_file(path)
        return False

    def _serve_index(self):
        filepath = os.path.join(TEMPLATE_DIR, "index.html")
        return self._serve_file_path(filepath)

    def _serve_file(self, path):
        filepath = os.path.join(os.path.dirname(os.path.abspath(__file__)), path.lstrip("/"))
        return self._serve_file_path(filepath)

    def _serve_file_path(self, filepath):
        if not os.path.exists(filepath):
            self._send_error_json("Not Found", 404)
            return True
        ext = os.path.splitext(filepath)[1].lower()
        mime = MIME_TYPES.get(ext, "application/octet-stream")
        self.send_response(200)
        self.send_header("Content-Type", mime)
        self.end_headers()
        with open(filepath, "rb") as f:
            self.wfile.write(f.read())
        return True

    # ===== API 路由 =====
    def _api_routes(self):
        path = self.path.split("?")[0]

        # 统计
        if path == "/api/stats" and self.command == "GET":
            self._handle_stats(); return True

        # 药材
        if path == "/api/herbs/all" and self.command == "GET":
            self._handle_all_herbs(); return True
        if path == "/api/herbs" and self.command == "GET":
            self._handle_list_herbs(); return True
        if path == "/api/herbs" and self.command == "POST":
            self._handle_create_herb(); return True
        if path.startswith("/api/herbs/") and self.command == "GET":
            self._handle_get_herb(path); return True
        if path.startswith("/api/herbs/") and self.command == "PUT":
            self._handle_update_herb(path); return True
        if path.startswith("/api/herbs/") and self.command == "DELETE":
            self._handle_delete_herb(path); return True

        # 方剂
        if path == "/api/prescriptions" and self.command == "GET":
            self._handle_list_prescriptions(); return True
        if path.startswith("/api/prescriptions/") and self.command == "GET":
            self._handle_get_prescription(path); return True

        # 药理学 (只读)
        if path == "/api/pharmacology" and self.command == "GET":
            self._handle_list_pharmacology(); return True
        if path.startswith("/api/pharmacology/") and self.command == "GET":
            self._handle_get_pharmacology(path); return True

        # 成分
        if path == "/api/components" and self.command == "GET":
            self._handle_list_components(); return True
        if path == "/api/components" and self.command == "POST":
            self._handle_create_component(); return True
        if path.startswith("/api/components/") and self.command == "GET":
            self._handle_get_component(path); return True
        if path.startswith("/api/components/") and self.command == "PUT":
            self._handle_update_component(path); return True
        if path.startswith("/api/components/") and self.command == "DELETE":
            self._handle_delete_component(path); return True

        # 研究
        if path == "/api/studies" and self.command == "GET":
            self._handle_list_studies(); return True
        if path == "/api/studies" and self.command == "POST":
            self._handle_create_study(); return True
        if path.startswith("/api/studies/") and self.command == "GET":
            self._handle_get_study(path); return True
        if path.startswith("/api/studies/") and self.command == "PUT":
            self._handle_update_study(path); return True
        if path.startswith("/api/studies/") and self.command == "DELETE":
            self._handle_delete_study(path); return True

        # 导入
        if path == "/api/import/preview" and self.command == "POST":
            self._handle_import_preview(); return True
        if path == "/api/import/confirm" and self.command == "POST":
            self._handle_import_confirm(); return True

        # 图谱
        if path == "/api/graph/query" and self.command == "GET":
            self._handle_graph_query(); return True
        if path == "/api/graph/stats" and self.command == "GET":
            self._handle_graph_stats(); return True
        if path == "/api/neo4j/sync" and self.command == "POST":
            self._handle_neo4j_sync(); return True
        if path == "/api/neo4j/health" and self.command == "GET":
            self._handle_neo4j_health(); return True

        # 知识库
        if path == "/api/knowledge/search" and self.command == "GET":
            self._handle_knowledge_search(); return True
        if path == "/api/knowledge/rebuild" and self.command == "POST":
            self._handle_knowledge_rebuild(); return True

        # 问答
        if path == "/api/qa/models" and self.command == "GET":
            self._handle_qa_models(); return True
        if path == "/api/qa/search" and self.command == "POST":
            self._handle_qa_search(); return True
        if path == "/api/qa/history" and self.command == "GET":
            self._handle_qa_history(); return True
        if path == "/api/qa/rate" and self.command == "POST":
            self._handle_qa_rate(); return True

        # 病症
        if path == "/api/diseases" and self.command == "GET":
            self._handle_list_diseases(); return True
        if path == "/api/diseases" and self.command == "POST":
            self._handle_create_disease(); return True
        if path.startswith("/api/diseases/") and self.command == "GET":
            self._handle_get_disease(path); return True
        if path.startswith("/api/diseases/") and self.command == "PUT":
            self._handle_update_disease(path); return True
        if path.startswith("/api/diseases/") and self.command == "DELETE":
            self._handle_delete_disease(path); return True

        return False

    def _get_id_from_path(self, path):
        parts = path.rstrip("/").split("/")
        last = parts[-1] if parts else ""
        try:
            return int(last)
        except (ValueError, IndexError):
            return last if last else None

    # ===== 统计 =====
    def _handle_stats(self):
        conn = get_db()
        herb_count = conn.execute("SELECT COUNT(*) FROM herb").fetchone()[0]
        pres_count = conn.execute("SELECT COUNT(*) FROM prescription").fetchone()[0]
        comp_count = conn.execute("SELECT COUNT(*) FROM chemical_component").fetchone()[0]
        disease_count = conn.execute("SELECT COUNT(*) FROM disease").fetchone()[0]
        pharm_count = conn.execute("SELECT COUNT(*) FROM pharmacology").fetchone()[0]
        herb_disease_count = conn.execute("SELECT COUNT(*) FROM herb_disease").fetchone()[0]
        total_relations = herb_disease_count + conn.execute("SELECT COUNT(*) FROM prescription_herb").fetchone()[0]

        recent = conn.execute("""
            SELECT 'herb' AS type, name AS title, created_at FROM herb
            UNION ALL SELECT 'prescription', name, created_at FROM prescription
            UNION ALL SELECT 'component', name, created_at FROM chemical_component
            UNION ALL SELECT 'disease', name, created_at FROM disease
            ORDER BY created_at DESC LIMIT 10
        """).fetchall()

        taste_dist = conn.execute("""
            SELECT taste, COUNT(*) AS cnt FROM herb WHERE taste != '' GROUP BY taste ORDER BY cnt DESC
        """).fetchall()

        conn.close()
        self._send_json({"success": True, "data": {
            "counts": {"herbs": herb_count, "prescriptions": pres_count, "components": comp_count, "diseases": disease_count, "pharmacology": pharm_count, "relations": total_relations},
            "recent": [dict(r) for r in recent],
            "taste_distribution": [dict(t) for t in taste_dist]
        }})

    # ===== CRUD 通用辅助 =====
    def _parse_list_args(self):
        params = self._parse_query()
        return {
            "search": self._get_query_param(params, "search"),
            "page": int(self._get_query_param(params, "page", "1")),
            "page_size": int(self._get_query_param(params, "page_size", "10")),
        }

    def _build_where(self, conditions, params_list):
        if not conditions:
            return "", params_list
        return " WHERE " + " AND ".join(conditions), params_list

    def _list_response(self, conn, table, fields, conditions, params, order="id DESC", page=1, page_size=10):
        where, params = self._build_where(conditions, params)
        total = conn.execute("SELECT COUNT(*) FROM " + table + where, params).fetchone()[0]
        offset = (page - 1) * page_size
        rows = conn.execute(
            "SELECT * FROM " + table + where + " ORDER BY " + order + " LIMIT ? OFFSET ?",
            params + [page_size, offset]
        ).fetchall()
        return {
            "items": [dict(r) for r in rows],
            "total": total,
            "page": page,
            "page_size": page_size
        }

    # ===== 药材 API =====
    def _handle_all_herbs(self):
        conn = get_db()
        rows = conn.execute("SELECT tcmbank_id AS id, name FROM herb ORDER BY name").fetchall()
        conn.close()
        self._send_json({"success": True, "data": [dict(r) for r in rows]})

    def _handle_list_herbs(self):
        args = self._parse_list_args()
        params = self._parse_query()
        conditions = []
        vals = []

        if args["search"]:
            conditions.append("(name LIKE ? OR latin_name LIKE ? OR efficacy LIKE ? OR pinyin_name LIKE ? OR tcm_name_en LIKE ? OR indication LIKE ?)")
            like = "%" + args["search"] + "%"
            vals.extend([like, like, like, like, like, like])
        nature = self._get_query_param(params, "nature")
        if nature:
            conditions.append("nature = ?")
            vals.append(nature)
        taste = self._get_query_param(params, "taste")
        if taste:
            conditions.append("taste LIKE ?")
            vals.append("%" + taste + "%")
        meridian = self._get_query_param(params, "meridian")
        if meridian:
            conditions.append("meridian LIKE ?")
            vals.append("%" + meridian + "%")

        conn = get_db()
        result = self._list_response(conn, "herb", "*", conditions, vals, page=args["page"], page_size=args["page_size"], order="tcmbank_id DESC")
        conn.close()
        self._send_json({"success": True, **result})

    def _handle_get_herb(self, path):
        herb_id = self._get_id_from_path(path)
        if herb_id is None:
            return self._send_error_json("无效ID", 400)
        conn = get_db()
        herb = conn.execute("SELECT * FROM herb WHERE tcmbank_id = ?", (herb_id,)).fetchone()
        if not herb:
            conn.close()
            return self._send_error_json("药材不存在", 404)
        result = dict(herb)
        pres = conn.execute("""
            SELECT p.cloud_id, p.name, ph.dosage FROM prescription_herb ph
            JOIN prescription p ON ph.prescription_id = p.cloud_id WHERE ph.herb_id = ?
        """, (herb_id,)).fetchall()
        result["prescriptions"] = [dict(p) for p in pres]
        comps = conn.execute("SELECT * FROM chemical_component WHERE herb_id = ?", (herb_id,)).fetchall()
        result["components"] = [dict(c) for c in comps]
        conn.close()
        self._send_json({"success": True, "data": result})

    def _handle_create_herb(self):
        data = json.loads(self._read_body())
        if not data.get("name"):
            return self._send_error_json("药材名称不能为空", 400)
        conn = get_db()
        cur = conn.execute("""
            INSERT INTO herb (name, latin_name, category, nature, taste, meridian, efficacy, toxicity, dosage, description,
                tcmbank_id, level1_name_en, pinyin_name, tcm_name_en, use_part, indication,
                clinical_manifestations, therapeutic_en_class, therapeutic_cn_class,
                tcmid_id, tcm_id_id, symmap_id, tcmsp_id, herb_external_id)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """, (data.get("name",""), data.get("latin_name",""), data.get("category",""), data.get("nature",""),
              data.get("taste",""), data.get("meridian",""), data.get("efficacy",""), data.get("toxicity",""),
              data.get("dosage",""), data.get("description",""),
              data.get("tcmbank_id",""), data.get("level1_name_en",""), data.get("pinyin_name",""),
              data.get("tcm_name_en",""), data.get("use_part",""), data.get("indication",""),
              data.get("clinical_manifestations",""), data.get("therapeutic_en_class",""),
              data.get("therapeutic_cn_class",""), data.get("tcmid_id",""), data.get("tcm_id_id",""),
              data.get("symmap_id",""), data.get("tcmsp_id",""), data.get("herb_external_id","")))
        conn.commit()
        hid = cur.lastrowid
        conn.close()
        self._send_json({"success": True, "data": {"id": hid}}, 201)

    def _handle_update_herb(self, path):
        herb_id = self._get_id_from_path(path)
        if herb_id is None:
            return self._send_error_json("无效ID", 400)
        data = json.loads(self._read_body())
        conn = get_db()
        conn.execute("""
            UPDATE herb SET name=?, latin_name=?, category=?, nature=?, taste=?, meridian=?,
            efficacy=?, toxicity=?, dosage=?, description=?,
            tcmbank_id=?, level1_name_en=?, pinyin_name=?, tcm_name_en=?, use_part=?, indication=?,
            clinical_manifestations=?, therapeutic_en_class=?, therapeutic_cn_class=?,
            tcmid_id=?, tcm_id_id=?, symmap_id=?, tcmsp_id=?, herb_external_id=?,
            updated_at=CURRENT_TIMESTAMP WHERE tcmbank_id=?
        """, (data.get("name",""), data.get("latin_name",""), data.get("category",""), data.get("nature",""),
              data.get("taste",""), data.get("meridian",""), data.get("efficacy",""), data.get("toxicity",""),
              data.get("dosage",""), data.get("description",""),
              data.get("tcmbank_id",""), data.get("level1_name_en",""), data.get("pinyin_name",""),
              data.get("tcm_name_en",""), data.get("use_part",""), data.get("indication",""),
              data.get("clinical_manifestations",""), data.get("therapeutic_en_class",""),
              data.get("therapeutic_cn_class",""), data.get("tcmid_id",""), data.get("tcm_id_id",""),
              data.get("symmap_id",""), data.get("tcmsp_id",""), data.get("herb_external_id",""),
              herb_id))
        conn.commit()
        conn.close()
        self._send_json({"success": True, "data": {"id": herb_id}})

    def _handle_delete_herb(self, path):
        herb_id = self._get_id_from_path(path)
        if herb_id is None:
            return self._send_error_json("无效ID", 400)
        conn = get_db()
        conn.execute("DELETE FROM herb WHERE tcmbank_id = ?", (herb_id,))
        conn.commit()
        conn.close()
        self._send_json({"success": True, "data": {"id": herb_id}})

    # ===== 方剂 API (只读，数据来自CSV导入) =====
    def _handle_list_prescriptions(self):
        args = self._parse_list_args()
        params = self._parse_query()
        conditions = []
        vals = []

        if args["search"]:
            conditions.append("(name LIKE ? OR efficacy LIKE ?)")
            like = "%" + args["search"] + "%"
            vals.extend([like, like])
        efficacy = self._get_query_param(params, "efficacy")
        if efficacy:
            conditions.append("efficacy LIKE ?")
            vals.append("%" + efficacy + "%")

        conn = get_db()
        where, v = self._build_where(conditions, vals)
        total = conn.execute("SELECT COUNT(*) FROM prescription" + where, v).fetchone()[0]
        offset = (args["page"] - 1) * args["page_size"]
        rows = conn.execute(
            "SELECT * FROM prescription" + where + " ORDER BY cloud_id DESC LIMIT ? OFFSET ?",
            v + [args["page_size"], offset]
        ).fetchall()

        items = []
        for r in rows:
            item = dict(r)
            item["id"] = r["cloud_id"]
            herbs = conn.execute("""
                SELECT h.tcmbank_id AS id, h.name, ph.dosage FROM prescription_herb ph
                JOIN herb h ON ph.herb_id = h.tcmbank_id WHERE ph.prescription_id = ?
            """, (r["cloud_id"],)).fetchall()
            item["herbs"] = [dict(h) for h in herbs]
            items.append(item)

        conn.close()
        self._send_json({"success": True, "items": items, "total": total, "page": args["page"], "page_size": args["page_size"]})

    def _handle_get_prescription(self, path):
        pres_id = self._get_id_from_path(path)
        if pres_id is None:
            return self._send_error_json("无效ID", 400)
        conn = get_db()
        pres = conn.execute("SELECT * FROM prescription WHERE cloud_id = ?", (pres_id,)).fetchone()
        if not pres:
            conn.close()
            return self._send_error_json("方剂不存在", 404)
        result = dict(pres)
        result["id"] = result["cloud_id"]
        herbs = conn.execute("""
            SELECT h.tcmbank_id AS id, h.name, ph.dosage FROM prescription_herb ph
            JOIN herb h ON ph.herb_id = h.tcmbank_id WHERE ph.prescription_id = ?
        """, (pres_id,)).fetchall()
        result["herbs"] = [dict(h) for h in herbs]
        conn.close()
        self._send_json({"success": True, "data": result})

    # ===== 药理学 API (只读) =====
    def _handle_list_pharmacology(self):
        args = self._parse_list_args()
        params = self._parse_query()
        conditions = []
        vals = []

        if args["search"]:
            conditions.append("name LIKE ?")
            vals.append("%" + args["search"] + "%")

        conn = get_db()
        where, v = self._build_where(conditions, vals)
        total = conn.execute("SELECT COUNT(*) FROM pharmacology" + where, v).fetchone()[0]
        offset = (args["page"] - 1) * args["page_size"]
        rows = conn.execute(
            "SELECT * FROM pharmacology" + where + " ORDER BY cloud_id DESC LIMIT ? OFFSET ?",
            v + [args["page_size"], offset]
        ).fetchall()
        conn.close()
        items = []
        for r in rows:
            d = dict(r)
            d["id"] = d["cloud_id"]
            items.append(d)
        self._send_json({"success": True, "items": items, "total": total, "page": args["page"], "page_size": args["page_size"]})

    def _handle_get_pharmacology(self, path):
        pharm_id = self._get_id_from_path(path)
        if pharm_id is None:
            return self._send_error_json("无效ID", 400)
        conn = get_db()
        pharm = conn.execute("SELECT * FROM pharmacology WHERE cloud_id = ?", (pharm_id,)).fetchone()
        if not pharm:
            conn.close()
            return self._send_error_json("药理学项目不存在", 404)
        result = dict(pharm)
        result["id"] = result["cloud_id"]
        herbs = conn.execute("""
            SELECT h.tcmbank_id AS id, h.name FROM herb_pharmacology hp
            JOIN herb h ON hp.herb_id = h.tcmbank_id WHERE hp.pharmacology_id = ?
        """, (pharm_id,)).fetchall()
        result["herbs"] = [dict(h) for h in herbs]
        ingredients = conn.execute("""
            SELECT c.cloud_id AS id, c.name FROM ingredient_pharmacology ip
            JOIN chemical_component c ON ip.ingredient_id = c.cloud_id WHERE ip.pharmacology_id = ?
        """, (pharm_id,)).fetchall()
        result["ingredients"] = [dict(i) for i in ingredients]
        conn.close()
        self._send_json({"success": True, "data": result})

    # ===== 成分 API =====
    def _handle_list_components(self):
        args = self._parse_list_args()
        params = self._parse_query()
        conditions = []
        vals = []

        if args["search"]:
            conditions.append("(c.name LIKE ? OR c.cas_number LIKE ?)")
            like = "%" + args["search"] + "%"
            vals.extend([like, like])
        formula = self._get_query_param(params, "formula")
        if formula:
            conditions.append("c.formula LIKE ?")
            vals.append("%" + formula + "%")
        herb_id = self._get_query_param(params, "herb_id")
        if herb_id:
            conditions.append("c.herb_id = ?")
            vals.append(herb_id)

        conn = get_db()
        where, v = self._build_where(conditions, vals)
        total = conn.execute("SELECT COUNT(*) FROM chemical_component c" + where, v).fetchone()[0]
        offset = (args["page"] - 1) * args["page_size"]
        rows = conn.execute("""
            SELECT c.*, h.name AS herb_name FROM chemical_component c
            LEFT JOIN herb h ON c.herb_id = h.tcmbank_id
            """ + where + " ORDER BY c.cloud_id DESC LIMIT ? OFFSET ?",
            v + [args["page_size"], offset]
        ).fetchall()
        conn.close()
        self._send_json({"success": True, "items": [dict(r) for r in rows], "total": total, "page": args["page"], "page_size": args["page_size"]})

    def _handle_get_component(self, path):
        comp_id = self._get_id_from_path(path)
        if comp_id is None:
            return self._send_error_json("无效ID", 400)
        conn = get_db()
        comp = conn.execute("""
            SELECT c.*, h.name AS herb_name FROM chemical_component c
            LEFT JOIN herb h ON c.herb_id = h.tcmbank_id WHERE c.cloud_id = ?
        """, (comp_id,)).fetchone()
        conn.close()
        if not comp:
            return self._send_error_json("成分不存在", 404)
        self._send_json({"success": True, "data": dict(comp)})

    def _handle_create_component(self):
        data = json.loads(self._read_body())
        if not data.get("name"):
            return self._send_error_json("成分名称不能为空", 400)
        conn = get_db()
        cur = conn.execute("""
            INSERT INTO chemical_component (name, formula, cas_number, herb_id, bioactivity)
            VALUES (?, ?, ?, ?, ?)
        """, (data.get("name",""), data.get("formula",""), data.get("cas_number",""),
              data.get("herb_id") or None, data.get("bioactivity","")))
        conn.commit()
        cid = cur.lastrowid
        conn.close()
        self._send_json({"success": True, "data": {"id": cid}}, 201)

    def _handle_update_component(self, path):
        comp_id = self._get_id_from_path(path)
        if comp_id is None:
            return self._send_error_json("无效ID", 400)
        data = json.loads(self._read_body())
        conn = get_db()
        conn.execute("""
            UPDATE chemical_component SET name=?, formula=?, cas_number=?, herb_id=?, bioactivity=? WHERE id=?
        """, (data.get("name",""), data.get("formula",""), data.get("cas_number",""),
              data.get("herb_id") or None, data.get("bioactivity",""), comp_id))
        conn.commit()
        conn.close()
        self._send_json({"success": True, "data": {"id": comp_id}})

    def _handle_delete_component(self, path):
        comp_id = self._get_id_from_path(path)
        if comp_id is None:
            return self._send_error_json("无效ID", 400)
        conn = get_db()
        conn.execute("DELETE FROM chemical_component WHERE id = ?", (comp_id,))
        conn.commit()
        conn.close()
        self._send_json({"success": True, "data": {"id": comp_id}})

    # ===== 研究 API =====
    def _handle_list_studies(self):
        args = self._parse_list_args()
        params = self._parse_query()
        conditions = []
        vals = []

        if args["search"]:
            conditions.append("(s.title LIKE ? OR s.effect LIKE ? OR s.summary LIKE ?)")
            like = "%" + args["search"] + "%"
            vals.extend([like, like, like])
        herb_id = self._get_query_param(params, "herb_id")
        if herb_id:
            conditions.append("s.herb_id = ?")
            vals.append(int(herb_id))
        component_id = self._get_query_param(params, "component_id")
        if component_id:
            conditions.append("s.component_id = ?")
            vals.append(int(component_id))

        conn = get_db()
        where, v = self._build_where(conditions, vals)
        total = conn.execute("SELECT COUNT(*) FROM pharma_study s" + where, v).fetchone()[0]
        offset = (args["page"] - 1) * args["page_size"]
        rows = conn.execute("""
            SELECT s.*, h.name AS herb_name, c.name AS component_name FROM pharma_study s
            LEFT JOIN herb h ON s.herb_id = h.id
            LEFT JOIN chemical_component c ON s.component_id = c.id
            """ + where + " ORDER BY s.id DESC LIMIT ? OFFSET ?",
            v + [args["page_size"], offset]
        ).fetchall()
        conn.close()
        self._send_json({"success": True, "items": [dict(r) for r in rows], "total": total, "page": args["page"], "page_size": args["page_size"]})

    def _handle_get_study(self, path):
        study_id = self._get_id_from_path(path)
        if study_id is None:
            return self._send_error_json("无效ID", 400)
        conn = get_db()
        study = conn.execute("""
            SELECT s.*, h.name AS herb_name, c.name AS component_name FROM pharma_study s
            LEFT JOIN herb h ON s.herb_id = h.id
            LEFT JOIN chemical_component c ON s.component_id = c.id WHERE s.id = ?
        """, (study_id,)).fetchone()
        conn.close()
        if not study:
            return self._send_error_json("研究不存在", 404)
        self._send_json({"success": True, "data": dict(study)})

    def _handle_create_study(self):
        data = json.loads(self._read_body())
        if not data.get("title"):
            return self._send_error_json("研究标题不能为空", 400)
        conn = get_db()
        cur = conn.execute("""
            INSERT INTO pharma_study (title, herb_id, component_id, effect, mechanism, reference, summary)
            VALUES (?, ?, ?, ?, ?, ?, ?)
        """, (data.get("title",""), data.get("herb_id") or None, data.get("component_id") or None,
              data.get("effect",""), data.get("mechanism",""), data.get("reference",""), data.get("summary","")))
        conn.commit()
        sid = cur.lastrowid
        conn.close()
        self._send_json({"success": True, "data": {"id": sid}}, 201)

    def _handle_update_study(self, path):
        study_id = self._get_id_from_path(path)
        if study_id is None:
            return self._send_error_json("无效ID", 400)
        data = json.loads(self._read_body())
        conn = get_db()
        conn.execute("""
            UPDATE pharma_study SET title=?, herb_id=?, component_id=?, effect=?, mechanism=?, reference=?, summary=?
            WHERE id=?
        """, (data.get("title",""), data.get("herb_id") or None, data.get("component_id") or None,
              data.get("effect",""), data.get("mechanism",""), data.get("reference",""), data.get("summary",""), study_id))
        conn.commit()
        conn.close()
        self._send_json({"success": True, "data": {"id": study_id}})

    def _handle_delete_study(self, path):
        study_id = self._get_id_from_path(path)
        if study_id is None:
            return self._send_error_json("无效ID", 400)
        conn = get_db()
        conn.execute("DELETE FROM pharma_study WHERE id = ?", (study_id,))
        conn.commit()
        conn.close()
        self._send_json({"success": True, "data": {"id": study_id}})

    # ===== 导入 API =====
    def _parse_upload_body(self):
        """解析 multipart/form-data"""
        content_type = self.headers.get("Content-Type", "")
        if "multipart/form-data" not in content_type:
            return None
        # 从 Content-Type 中提取 boundary
        boundary = None
        for part in content_type.split(";"):
            part = part.strip()
            if part.startswith("boundary="):
                boundary = part[9:].strip('"')
                break
        if not boundary:
            return None

        body = self._read_body()
        # 手动解析 multipart
        boundary_bytes = boundary.encode("utf-8")
        parts = body.split(b"--" + boundary_bytes)

        result = {"fields": {}, "files": {}}
        for part in parts[1:-1]:  # 跳过第一个空和最后一个 --
            if b"\r\n\r\n" not in part:
                continue
            header_section, content = part.split(b"\r\n\r\n", 1)
            content = content.rstrip(b"\r\n")  # 去掉尾部换行

            headers = {}
            for line in header_section.decode("utf-8", errors="ignore").split("\r\n"):
                if ":" in line:
                    key, val = line.split(":", 1)
                    headers[key.strip().lower()] = val.strip()

            disposition = headers.get("content-disposition", "")
            name = None
            filename = None
            for disp_part in disposition.split(";"):
                disp_part = disp_part.strip()
                if disp_part.startswith("name="):
                    name = disp_part[5:].strip('"')
                elif disp_part.startswith("filename="):
                    filename = disp_part[9:].strip('"')

            if filename:
                result["files"][name] = {"filename": filename, "content": content}
            elif name:
                result["fields"][name] = content.decode("utf-8", errors="ignore")

        return result

    def _handle_import_preview(self):
        parsed = self._parse_upload_body()
        if not parsed or "file" not in parsed["files"]:
            return self._send_error_json("未上传文件", 400)

        file_info = parsed["files"]["file"]
        file_type = parsed["fields"].get("type", "csv")
        filename = file_info["filename"]
        content = file_info["content"]

        ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
        valid = {"csv": ["csv"], "json": ["json"], "excel": ["xlsx", "xls"]}
        if ext not in valid.get(file_type, []):
            return self._send_error_json("文件扩展名与类型不匹配", 400)

        # 保存临时文件
        temp_name = str(uuid.uuid4()) + "." + ext
        filepath = os.path.join(UPLOAD_DIR, temp_name)
        with open(filepath, "wb") as f:
            f.write(content)

        try:
            if file_type == "csv":
                rows = self._parse_csv(filepath)
            elif file_type == "json":
                rows = self._parse_json(filepath)
            elif file_type == "excel":
                rows = self._parse_excel(filepath)
            else:
                return self._send_error_json("不支持的文件类型", 400)

            headers = list(rows[0].keys()) if rows else []
            self._send_json({"success": True, "data": {
                "headers": headers,
                "rows": rows[:20],
                "total": len(rows),
                "temp_file": temp_name
            }})
        except Exception as e:
            self._send_error_json("文件解析失败: " + str(e), 400)

    def _handle_import_confirm(self):
        data = json.loads(self._read_body())
        temp_file = data.get("temp_file")
        file_type = data.get("type", "csv")
        entity_type = data.get("entity_type")

        if not temp_file or not entity_type:
            return self._send_error_json("参数不完整", 400)

        filepath = os.path.join(UPLOAD_DIR, temp_file)
        if not os.path.exists(filepath):
            return self._send_error_json("临时文件不存在", 400)

        try:
            if file_type == "csv":
                rows = self._parse_csv(filepath)
            elif file_type == "json":
                rows = self._parse_json(filepath)
            elif file_type == "excel":
                rows = self._parse_excel(filepath)
            else:
                return self._send_error_json("不支持的文件类型", 400)

            table_map = {"herb": "herb", "prescription": "prescription", "component": "chemical_component", "study": "pharma_study"}
            field_map = {
                "herb": ["name", "latin_name", "category", "nature", "taste", "meridian", "efficacy", "toxicity", "dosage", "description",
                         "tcmbank_id", "level1_name_en", "pinyin_name", "tcm_name_en", "use_part", "indication",
                         "clinical_manifestations", "therapeutic_en_class", "therapeutic_cn_class",
                         "tcmid_id", "tcm_id_id", "symmap_id", "tcmsp_id", "herb_external_id"],
                "prescription": ["name", "category", "efficacy", "indications", "source", "description"],
                "component": ["name", "formula", "cas_number", "herb_id", "bioactivity"],
                "study": ["title", "herb_id", "component_id", "effect", "mechanism", "reference", "summary"],
            }

            if entity_type == "tcmbank":
                result = self._import_tcmbank(rows)
                os.remove(filepath)
                return result

            table = table_map[entity_type]
            fields = field_map[entity_type]
            conn = get_db()
            count = 0
            error_count = 0
            for row in rows:
                try:
                    values = [(str(row.get(f, "") or "")) for f in fields]
                    placeholders = ", ".join(["?"] * len(fields))
                    cols = ", ".join(fields)
                    conn.execute("INSERT INTO " + table + " (" + cols + ") VALUES (" + placeholders + ")", values)
                    count += 1
                except Exception:
                    error_count += 1
                    continue
            conn.commit()
            conn.close()
            os.remove(filepath)
            self._send_json({"success": True, "data": {"imported": count, "errors": error_count, "entity_type": entity_type}})
        except Exception as e:
            self._send_error_json("导入失败: " + str(e), 400)

    def _import_tcmbank(self, rows):
        """导入 TCMBank 格式的药材数据"""
        conn = get_db()
        count = 0
        for row in rows:
            # Properties 字段拆分: "Warm;Acrid" → nature / taste
            props = row.get("Properties", "") or ""
            parts = [p.strip() for p in props.split(";") if p.strip()]
            nature = parts[0] if len(parts) >= 1 else ""
            taste = parts[1] if len(parts) >= 2 else ""

            name = row.get("TCM_name", "") or row.get("TCM_name_en", "")

            conn.execute("""
                INSERT INTO herb (name, latin_name, category, nature, taste, meridian, efficacy, toxicity, dosage,
                    tcmbank_id, level1_name_en, pinyin_name, tcm_name_en, use_part, indication,
                    clinical_manifestations, therapeutic_en_class, therapeutic_cn_class,
                    tcmid_id, tcm_id_id, symmap_id, tcmsp_id, herb_external_id)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """, (
                name or "",
                row.get("Herb_latin_name", "") or "",
                row.get("level2_name", "") or "",
                nature,
                taste,
                row.get("Meridians", "") or "",
                row.get("Function", "") or "",
                row.get("Toxicity", "") or "",
                "",
                row.get("TCMBank_ID", "") or "",
                row.get("level1_name_en", "") or "",
                row.get("Herb_pinyin_name", "") or "",
                row.get("TCM_name_en", "") or "",
                row.get("UsePart", "") or "",
                row.get("Indication", "") or "",
                row.get("Clinical_manifestations", "") or "",
                row.get("Therapeutic_en_class", "") or "",
                row.get("Therapeutic_cn_class", "") or "",
                row.get("TCMID_id", "") or "",
                row.get("TCM_ID_id", "") or "",
                row.get("SymMap_id", "") or "",
                row.get("TCMSP_id", "") or "",
                row.get("Herb_ID", "") or "",
            ))
            count += 1
        conn.commit()
        conn.close()
        self._send_json({"success": True, "data": {"imported": count, "entity_type": "tcmbank"}})

    def _parse_csv(self, filepath):
        rows = []
        with open(filepath, "r", encoding="utf-8-sig") as f:
            reader = csv.DictReader(f)
            for row in reader:
                rows.append(row)
        return rows

    def _parse_json(self, filepath):
        with open(filepath, "r", encoding="utf-8") as f:
            data = json.load(f)
        if isinstance(data, list):
            return data
        if isinstance(data, dict) and "data" in data:
            return data["data"]
        return [data]

    def _parse_excel(self, filepath):
        try:
            import openpyxl
        except ImportError:
            raise ImportError("Excel 导入需要 openpyxl 库，请运行: pip install openpyxl")
        wb = openpyxl.load_workbook(filepath, read_only=True)
        ws = wb.active
        rows_iter = ws.iter_rows(values_only=True)
        headers = [str(h) if h else "" for h in next(rows_iter)]
        result = []
        for row in rows_iter:
            result.append(dict(zip(headers, row)))
        wb.close()
        return result

    # ===== 主请求分发 =====
    def do_GET(self):
        if self._serve_static():
            return
        if self._api_routes():
            return
        self._send_error_json("Not Found", 404)

    def do_POST(self):
        if self._api_routes():
            return
        self._send_error_json("Not Found", 404)

    def do_PUT(self):
        if self._api_routes():
            return
        self._send_error_json("Not Found", 404)

    def do_DELETE(self):
        if self._api_routes():
            return
        self._send_error_json("Not Found", 404)


# ===== 图谱 API =====
    def _handle_graph_query(self):
        params = self._parse_query()
        entity_type = self._get_query_param(params, "entity_type")
        entity_id = self._get_query_param(params, "entity_id")
        if not entity_type or not entity_id:
            return self._send_error_json("缺少 entity_type 或 entity_id", 400)
        # entity_id 现在是文本格式（tcmbank_id/cloud_id）
        # 优先使用 Neo4j，失败则 fallback SQLite
        try:
            result = query_entity_graph(entity_type, entity_id)
        except Exception:
            result = GraphService.query_entity_graph(entity_type, entity_id)
        self._send_json({"success": True, "data": result})

    def _handle_graph_stats(self):
        # 优先使用 Neo4j，失败则 fallback SQLite
        try:
            stats = get_statistics()
        except Exception:
            stats = GraphService.get_statistics()
        self._send_json({"success": True, "data": stats})

    def _handle_neo4j_sync(self):
        """同步 SQLite 数据到 Neo4j"""
        try:
            stats = sync_all_to_neo4j()
            self._send_json({"success": True, "data": stats})
        except Exception as e:
            self._send_error_json("Neo4j 同步失败: " + str(e), 500)

    def _handle_neo4j_health(self):
        """检查 Neo4j 连接状态"""
        result = check_health()
        self._send_json({"success": True, "data": result})

    # ===== 知识库 API =====
    def _handle_knowledge_search(self):
        params = self._parse_query()
        query = self._get_query_param(params, "q")
        entity_type = self._get_query_param(params, "entity_type")
        if not query:
            return self._send_error_json("缺少搜索关键词 q", 400)
        results = KnowledgeService.search_chunks(query, entity_type=entity_type or None)
        self._send_json({"success": True, "data": results, "total": len(results)})

    def _handle_knowledge_rebuild(self):
        try:
            count = KnowledgeService.rebuild_all_entity_chunks()
            self._send_json({"success": True, "data": {"chunks": count}})
        except Exception as e:
            self._send_error_json("重建失败: " + str(e), 500)

    # ===== 问答 API =====
    def _handle_qa_models(self):
        """获取可用模型列表"""
        try:
            models = QAService.fetch_models()
            self._send_json({"success": True, "data": models})
        except Exception as e:
            self._send_error_json("获取模型列表失败: " + str(e), 500)

    def _handle_qa_search(self):
        data = json.loads(self._read_body())
        query = data.get("question", "")
        model = data.get("model", "agnes-2.0-flash")
        if not query:
            return self._send_error_json("问题不能为空", 400)
        result = QAService.qa_search(query, model=model)
        QAService.save_qa(query, result["answer"], result["context"], result["model"])
        self._send_json({"success": True, "data": result})

    def _handle_qa_history(self):
        rows = QAService.get_recent_qa()
        self._send_json({"success": True, "data": rows})

    def _handle_qa_rate(self):
        data = json.loads(self._read_body())
        qa_id = data.get("qa_id")
        rating = data.get("rating")
        if not qa_id or rating is None:
            return self._send_error_json("参数不完整", 400)
        QAService.rate_qa(qa_id, rating, data.get("feedback"))
        self._send_json({"success": True})

    # ===== 病症 CRUD API =====
    def _handle_list_diseases(self):
        args = self._parse_list_args()
        conn = get_db()
        conditions = []
        vals = []
        if args["search"]:
            conditions.append("(name LIKE ? OR category LIKE ?)")
            like = "%" + args["search"] + "%"
            vals.extend([like, like])
        where, v = self._build_where(conditions, vals)
        total = conn.execute("SELECT COUNT(*) FROM disease" + where, v).fetchone()[0]
        offset = (args["page"] - 1) * args["page_size"]
        rows = conn.execute(
            "SELECT * FROM disease" + where + " ORDER BY cloud_id DESC LIMIT ? OFFSET ?",
            v + [args["page_size"], offset]
        ).fetchall()
        conn.close()
        self._send_json({"success": True, "items": [dict(r) for r in rows], "total": total, "page": args["page"], "page_size": args["page_size"]})

    def _handle_get_disease(self, path):
        disease_id = self._get_id_from_path(path)
        if disease_id is None:
            return self._send_error_json("无效ID", 400)
        conn = get_db()
        disease = conn.execute("SELECT * FROM disease WHERE cloud_id = ?", (disease_id,)).fetchone()
        if not disease:
            conn.close()
            return self._send_error_json("病症不存在", 404)
        result = dict(disease)
        # 关联药材
        herbs = conn.execute("""
            SELECT h.tcmbank_id, h.name, hd.relationship_type, hd.evidence_level
            FROM herb_disease hd JOIN herb h ON hd.herb_tcmbank_id = h.tcmbank_id WHERE hd.disease_cloud_id = ?
        """, (disease_id,)).fetchall()
        result["herbs"] = [dict(h) for h in herbs]
        # 关联方剂
        pres = conn.execute("""
            SELECT p.cloud_id, p.name, pd.relationship_type, pd.evidence_level
            FROM prescription_disease pd JOIN prescription p ON pd.prescription_id = p.cloud_id WHERE pd.disease_id = ?
        """, (disease_id,)).fetchall()
        result["prescriptions"] = [dict(p) for p in pres]
        conn.close()
        self._send_json({"success": True, "data": result})

    def _handle_create_disease(self):
        data = json.loads(self._read_body())
        if not data.get("name"):
            return self._send_error_json("病症名称不能为空", 400)
        conn = get_db()
        cur = conn.execute("""
            INSERT INTO disease (name, category, description, tcm_syndrome) VALUES (?, ?, ?, ?)
        """, (data.get("name",""), data.get("category",""), data.get("description",""), data.get("tcm_syndrome","")))
        conn.commit()
        did = cur.lastrowid
        conn.close()
        self._send_json({"success": True, "data": {"id": did}}, 201)

    def _handle_update_disease(self, path):
        disease_id = self._get_id_from_path(path)
        if disease_id is None:
            return self._send_error_json("无效ID", 400)
        data = json.loads(self._read_body())
        conn = get_db()
        conn.execute("""
            UPDATE disease SET name=?, category=?, description=?, tcm_syndrome=? WHERE cloud_id=?
        """, (data.get("name",""), data.get("category",""), data.get("description",""), data.get("tcm_syndrome",""), disease_id))
        conn.commit()
        conn.close()
        self._send_json({"success": True, "data": {"id": disease_id}})

    def _handle_delete_disease(self, path):
        disease_id = self._get_id_from_path(path)
        if disease_id is None:
            return self._send_error_json("无效ID", 400)
        conn = get_db()
        conn.execute("DELETE FROM disease WHERE cloud_id = ?", (disease_id,))
        conn.commit()
        conn.close()
        self._send_json({"success": True, "data": {"id": disease_id}})


def main():
    print("main() 开始执行")
    os.makedirs(UPLOAD_DIR, exist_ok=True)
    # 重新初始化数据库（含新扩展表）
    init_db()
    print("init_db 完成")
    # 重建知识库分块
    try:
        chunk_count = KnowledgeService.rebuild_all_entity_chunks()
        print(f"  知识库分块已重建: {chunk_count} 个")
    except Exception as e:
        print(f"  知识库分块重建跳过: {e}")
    port = 5000
    server = HTTPServer(("0.0.0.0", port), TCMHandler)
    print("=" * 55)
    print("  中医药科学大数据管理平台 v2.0")
    print("  含 Neo4j 知识图谱 & LLM 问答系统扩展")
    print("=" * 55)
    print(f"  访问地址: http://127.0.0.1:{port}")
    print("  按 Ctrl+C 停止服务")
    print("  服务器准备就绪，等待请求...")
    sys.stdout.flush()
    try:
        server.serve_forever()
    except KeyboardInterrupt:
        print("\n服务已停止")
        server.server_close()


if __name__ == "__main__":
    try:
        main()
    except KeyboardInterrupt:
        print("\n服务已停止")
    except Exception as e:
        import traceback
        print(f"\n服务器启动失败: {e}")
        traceback.print_exc()