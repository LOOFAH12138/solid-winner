# -*- coding: utf-8 -*-
"""数据导入服务层"""
import csv
import json
import os
import io
from database import get_db


class ImportService:

    ENTITY_FIELDS = {
        "herb": ["tcmbank_id", "name", "latin_name", "category", "nature", "taste", "meridian", "pinyin_name", "tcm_name_en"],
        "disease": ["cloud_id", "name", "category", "description", "tcm_syndrome", "mesh_class"],
        "component": ["cloud_id", "name", "formula", "cas_number", "bioactivity"],
    }

    TABLE_MAP = {
        "herb": "herb",
        "disease": "disease",
        "component": "chemical_component",
    }

    @staticmethod
    def parse_csv(filepath):
        """解析CSV文件"""
        rows = []
        with open(filepath, "r", encoding="utf-8-sig") as f:
            reader = csv.DictReader(f)
            for row in reader:
                rows.append(row)
        return rows

    @staticmethod
    def parse_json(filepath):
        """解析JSON文件"""
        with open(filepath, "r", encoding="utf-8") as f:
            data = json.load(f)
        if isinstance(data, list):
            return data
        if isinstance(data, dict) and "data" in data:
            return data["data"]
        return [data]

    @staticmethod
    def parse_excel(filepath):
        """解析Excel文件"""
        try:
            import openpyxl
        except ImportError:
            raise ImportError("请安装 openpyxl: pip install openpyxl")
        wb = openpyxl.load_workbook(filepath, read_only=True)
        ws = wb.active
        rows_iter = ws.iter_rows(values_only=True)
        headers = [str(h) if h else "" for h in next(rows_iter)]
        result = []
        for row in rows_iter:
            result.append(dict(zip(headers, row)))
        wb.close()
        return result

    @classmethod
    def preview_file(cls, filepath, file_type):
        """预览文件内容"""
        if file_type == "csv":
            rows = cls.parse_csv(filepath)
        elif file_type == "json":
            rows = cls.parse_json(filepath)
        elif file_type == "excel":
            rows = cls.parse_excel(filepath)
        else:
            raise ValueError("不支持的文件类型: " + file_type)
        return {"headers": list(rows[0].keys()) if rows else [], "rows": rows[:20], "total": len(rows)}

    @classmethod
    def import_data(cls, filepath, file_type, entity_type):
        """将解析后的数据导入数据库"""
        if entity_type not in cls.ENTITY_FIELDS:
            raise ValueError("不支持的实体类型: " + entity_type)

        if file_type == "csv":
            rows = cls.parse_csv(filepath)
        elif file_type == "json":
            rows = cls.parse_json(filepath)
        elif file_type == "excel":
            rows = cls.parse_excel(filepath)
        else:
            raise ValueError("不支持的文件类型: " + file_type)

        table = cls.TABLE_MAP[entity_type]
        fields = cls.ENTITY_FIELDS[entity_type]
        conn = get_db()
        count = 0

        for row in rows:
            values = []
            for f in fields:
                val = row.get(f, "")
                if val is None:
                    val = ""
                values.append(val)
            placeholders = ", ".join(["?"] * len(fields))
            cols = ", ".join(fields)
            conn.execute(
                "INSERT INTO {} ({}) VALUES ({})".format(table, cols, placeholders),
                values
            )
            count += 1

        conn.commit()
        conn.close()
        return {"imported": count, "entity_type": entity_type}